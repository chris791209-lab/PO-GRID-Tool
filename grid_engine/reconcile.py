#!/usr/bin/env python3
"""
reconcile.py — pick the live PO version, expand assortments, cross-check the
customer master, and detect cross-numbered duplicate POs.

Version rule: DATE decides. The latest-dated copy of a PO number wins; earlier
copies are superseded and contribute nothing. Do not use `PO Type:` — SPS
reprints a revised PO still labelled "Original".

Date precedence:
  1. page-1 retrieval stamp   "2026/2/12 09:22 Fulfillment"
  2. filename suffix          "_0224.pdf"   (month >= 11 -> previous year)
  3. folder name              "12.24"

Same-number versioning is the easy half. The dangerous half is duplicates that
carry DIFFERENT PO numbers — see references/duplicate-pos.md. This script
fingerprints every PO by its (sku, qty) set and reports collisions rather than
resolving them, because only the account team knows which copy is live.

USAGE
  python reconcile.py --items work/items.json --outdir work \
      --master "Decor & Night Of=2026_C5_D240_Halloween_Decor.xlsx" \
      --master "Costumes=2026_C5_D240_Halloween_Costumes.xlsx" \
      --assortment "HW_26_FA_Box_Template.xlsx:Fixed Assortments" \
      --assortment "Halloween_2026_Assortment_Box_Proposal.xlsx:Assortment Proposal Details" \
      --exclude-po 4818185,6945142
"""
import argparse, json, os, re, warnings
from collections import defaultdict
from datetime import datetime

import openpyxl

warnings.filterwarnings('ignore')


# ---------------------------------------------------------------- versioning
QTY_TOL_PCT = 10.0   # standing rule (Chris, 2026-08-18): +-10% is normal variance


def version_date(h):
    s = h.get('retrieved') or ''
    m = re.match(r'(\d{4})/(\d{1,2})/(\d{1,2})\s+(\d{1,2}):(\d{2})', s)
    if m:
        y, mo, d, hh, mi = (int(x) for x in m.groups())
        return datetime(y, mo, d, hh, mi), 'stamp'
    fd = h.get('fname_date') or ''
    if len(fd) == 4:
        mo, d = int(fd[:2]), int(fd[2:])
        return datetime(2000 + (datetime.now().year % 100) - (1 if mo >= 11 else 0),
                        mo, d), 'filename'
    fol = re.match(r'(\d{2})\.(\d{2})', h.get('folder') or '')
    if fol:
        mo, d = int(fol.group(1)), int(fol.group(2))
        return datetime(2000 + (datetime.now().year % 100) - (1 if mo >= 11 else 0),
                        mo, d), 'folder'
    return datetime(1900, 1, 1), 'none'


def dedup(pos):
    groups = defaultdict(list)
    for h in pos:
        h['_dt'], h['_dtsrc'] = version_date(h)
        groups[(h['po'], h['dc'])].append(h)
    kept, superseded = [], []
    for key, hs in groups.items():
        hs.sort(key=lambda x: x['_dt'], reverse=True)
        win = hs[0]
        kept.append(win)
        for lose in hs[1:]:
            superseded.append({
                'po': key[0], 'dc': key[1],
                'dropped': lose['file'], 'kept': win['file'],
                'qty_changed': lose.get('total_qty') != win.get('total_qty'),
                'from_qty': lose.get('total_qty'), 'to_qty': win.get('total_qty'),
                'from_total': lose.get('po_total'), 'to_total': win.get('po_total'),
            })
    return kept, superseded


def fingerprint_duplicates(kept, items):
    """POs with an identical (sku, qty) set but different PO numbers."""
    per_file = defaultdict(list)
    for i in items:
        if i['kind'] in ('line', 'prepack'):
            per_file[i['file']].append((i['sku'], i['qty']))
    fp = defaultdict(list)
    for h in kept:
        key = tuple(sorted(per_file.get(h['file'], [])))
        if key:
            fp[key].append(h)
    out = []
    for key, hs in fp.items():
        if len({h['po'] for h in hs}) > 1:
            out.append({
                'skus': len(key), 'units': sum(q for _, q in key),
                'copies': [{'po': h['po'], 'dc': h['dc'], 'folder': h['folder'],
                            'po_total': h['po_total'], 'file': h['file'],
                            'date': h['_dt'].isoformat()} for h in sorted(hs, key=lambda x: x['_dt'])],
            })
    return out


# ------------------------------------------------------------------- masters
def load_master(path, sheet='Data'):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    hdr = [str(h).strip() if h is not None else '' for h in rows[0]]
    missing = []

    def col(*names, required=True):
        for n in names:
            if n in hdr:
                return hdr.index(n)
        if required:
            missing.append(names[0])
        return None

    ci = {
        'dpci': col('DPCI'),
        'desc': col('Product Description'),
        'plan': col('Ent Ttl Rcpt U'),
        'retail': col('Suggested Unit Retail'),
        'case': col('Case Unit Quantity'),
        'inner': col('Inner Pack Unit Quantity'),
        'fca': col('FCA Factory City Unit Cost'),
        'fob': col('FOB Unit Cost'),
        'factory': col('Factory Name'),
        'facid': col('Factory ID'),
        'style': col('Manufacturer Style # *', 'Manufacturer Style #'),
        'barcode': col('Barcode'),
        'hts': col('HTS Code', required=False),
        'coo': col('Import Country of Origin', required=False),
        'port': col('Port of Export', required=False),
        'subclass': col('Subclass Name', required=False),
        'material': col('Main Raw Material *', 'Primary Raw Material Type',
                        'Main Raw Material', required=False),
        'pack_format': col('Retail Packaging Format (1) *',
                           'Retail Packaging Format (1)', required=False),
    }
    if missing:
        raise SystemExit('master %s missing required column(s): %s\n'
                         'Locate columns by header text, never by position — '
                         'stop and ask rather than guessing.' % (path, missing))

    out = {}
    for r in rows[1:]:
        if not r or not r[ci['dpci']]:
            continue
        rec = {k: (r[i] if i is not None and i < len(r) else None) for k, i in ci.items()}
        rec['dpci'] = str(rec['dpci']).strip()
        rec['cost'] = rec['fca'] if rec['fca'] not in (None, '') else rec['fob']
        rec['cost_basis'] = ('FCA' if rec['fca'] not in (None, '')
                             else ('FOB' if rec['fob'] not in (None, '') else ''))
        out[rec['dpci']] = rec
    return out


DPCI_RE = re.compile(r'^\d{3}-\d{2}-\d{4}$')

# Column aliases — the FA Box Template and the Assortment Proposal name the same
# things differently. Match on header text, in this order of preference.
A_COMP = ('Component Item DPCI', 'Item DPCI')
A_ASST = ('Assortment DPCI',)
A_UNITS = ('Units in Assortment', '# Units in Assortment')
A_COST = ('Item Cost',)
A_BOX = ('FA box cost', 'Asst Cost')
A_STYLE = ('Vendor Asst Style #',)
A_DESC = ('Assortment Description',)


def load_assortments(specs):
    """specs: ["file.xlsx:Sheet Name", ...] — later files fill gaps, never overwrite."""
    asst = {}
    for spec in specs:
        path, _, sheet = spec.partition(':')
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        hi = None
        for i, r in enumerate(rows):
            cells = {str(c).strip() for c in r if c}
            if cells & set(A_COMP):
                hi = i
                break
        if hi is None:
            raise SystemExit('%s: no header row containing %s' % (path, ' / '.join(A_COMP)))
        hdr = [str(h).strip() if h else '' for h in rows[hi]]

        def pick(d, names):
            for n in names:
                if n in d and d[n] not in (None, ''):
                    return d[n]
            return None

        cur = None
        for r in rows[hi + 1:]:
            d = dict(zip(hdr, r))
            comp = str(pick(d, A_COMP) or '').strip()
            if not DPCI_RE.match(comp):
                continue
            a = str(pick(d, A_ASST) or '').strip()
            if a:
                cur = a
            a = a or cur
            if not a:
                continue
            e = asst.setdefault(a, {'style': None, 'desc': None, 'box_cost': None,
                                    'comp': [], 'src': os.path.basename(path)})
            for k, names in (('style', A_STYLE), ('desc', A_DESC)):
                v = pick(d, names)
                if v and not e[k]:
                    e[k] = str(v).strip()
            bc = pick(d, A_BOX)
            if bc and e['box_cost'] is None:
                e['box_cost'] = float(bc)
            if not any(c['dpci'] == comp for c in e['comp']):
                e['comp'].append({'dpci': comp,
                                  'units': int(pick(d, A_UNITS) or 0),
                                  'cost': float(pick(d, A_COST) or 0)})
    return asst


def fmt(sku):
    return '%s-%s-%s' % (sku[:3], sku[3:5], sku[5:])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--items', required=True)
    ap.add_argument('--master', action='append', required=True,
                    help='"Program Label=path.xlsx" (repeatable)')
    ap.add_argument('--assortment', action='append', default=[],
                    help='"path.xlsx:Sheet Name" (repeatable)')
    ap.add_argument('--exclude-po', default='',
                    help='comma-separated PO numbers to drop entirely')
    ap.add_argument('--qty-tolerance-pct', type=float, default=10.0,

                    help='treat |ordered-plan| within this %% of plan as normal (default 10; 0 = case-pack rule only)')

    ap.add_argument('--outdir', default='./work')
    a = ap.parse_args()
    global QTY_TOL_PCT
    QTY_TOL_PCT = a.qty_tolerance_pct
    os.makedirs(a.outdir, exist_ok=True)

    data = json.load(open(a.items, encoding='utf-8'))
    pos, items = data['pos'], data['items']
    excluded = {p.strip() for p in a.exclude_po.split(',') if p.strip()}
    pos = [h for h in pos if h['po'] not in excluded]

    kept, superseded = dedup(pos)
    keep_files = {h['file'] for h in kept}
    live = [i for i in items if i['file'] in keep_files]
    dup_groups = fingerprint_duplicates(kept, live)

    direct, embedded, ast_boxes = defaultdict(dict), defaultdict(dict), defaultdict(dict)
    for i in live:
        d, po = fmt(i['sku']), i['po']
        tgt = {'line': direct, 'ast': ast_boxes, 'prepack': embedded}.get(i['kind'])
        if tgt is not None:
            tgt[d][po] = tgt[d].get(po, 0) + i['qty']

    masters, all_master = {}, {}
    for spec in a.master:
        label, _, path = spec.partition('=')
        if not path:
            label, path = os.path.basename(spec), spec
        m = load_master(path)
        masters[label] = m
        for k, v in m.items():
            v['program'] = label
            all_master[k] = v

    asst = load_assortments(a.assortment) if a.assortment else {}

    ordered = set(direct) | set(embedded) | set(ast_boxes)
    findings = {'unknown_dpci': [], 'cost_mismatch': [], 'retail_mismatch': [],
                'qty_vs_plan': [], 'assortment_mismatch': [], 'not_ordered': [],
                'duplicate_po_groups': dup_groups}

    # assortment expansion cross-check
    for adpci, boxes_by_po in ast_boxes.items():
        spec = asst.get(adpci)
        if not spec:
            findings['assortment_mismatch'].append(
                {'assortment': adpci, 'issue': 'no definition in any supplied assortment file'})
            continue
        boxes = sum(boxes_by_po.values())
        if spec['box_cost']:
            calc = sum(c['units'] * c['cost'] for c in spec['comp'])
            if abs(calc - spec['box_cost']) > 0.01:
                findings['assortment_mismatch'].append(
                    {'assortment': adpci, 'issue': 'box cost %.4f != sum of components %.4f'
                                                   % (spec['box_cost'], calc)})
        for c in spec['comp']:
            want = boxes * c['units']
            got = sum(embedded.get(c['dpci'], {}).get(po, 0) for po in boxes_by_po)
            if want != got:
                findings['assortment_mismatch'].append({
                    'assortment': adpci, 'component': c['dpci'], 'boxes': boxes,
                    'units_per_box': c['units'], 'expected': want,
                    'in_po': got, 'diff': got - want})

    rows = []
    for d in sorted(ordered):
        if d in asst:
            continue                       # assortment master, not a stock DPCI
        m = all_master.get(d)
        dqty = sum(direct.get(d, {}).values())
        eqty = sum(embedded.get(d, {}).values())
        total = dqty + eqty
        if not m:
            findings['unknown_dpci'].append(
                {'dpci': d, 'ordered': total,
                 'pos': sorted(set(direct.get(d, {})) | set(embedded.get(d, {})))})
            continue

        plan = m['plan'] or 0
        case = m['case'] or 1
        diff = total - plan
        pct_off = abs(diff / plan * 100) if plan else 100.0
        if abs(diff) >= case and pct_off > QTY_TOL_PCT:
            findings['qty_vs_plan'].append({
                'dpci': d, 'program': m['program'], 'desc': str(m['desc'])[:45],
                'ordered': total, 'plan': round(plan), 'diff': round(diff),
                'case_pack': case, 'cases': round(diff / case, 1) if case else None,
                'pct': round(diff / plan * 100, 1) if plan else None})

        for i in live:
            if fmt(i['sku']) != d or i['kind'] not in ('line', 'prepack'):
                continue
            if i.get('unit') and m['cost'] not in (None, ''):
                if abs(float(i['unit']) - float(m['cost'])) > 0.005:
                    findings['cost_mismatch'].append(
                        {'dpci': d, 'po': i['po'], 'po_unit': i['unit'],
                         'master_cost': m['cost'], 'basis': m['cost_basis']})
            if i['kind'] == 'line' and i.get('resale') and m['retail'] not in (None, ''):
                if abs(float(i['resale']) - float(m['retail'])) > 0.005:
                    findings['retail_mismatch'].append(
                        {'dpci': d, 'po': i['po'], 'po_resale': i['resale'],
                         'master_retail': m['retail']})

        rows.append({
            'dpci': d, 'program': m['program'], 'desc': m['desc'], 'style': m['style'],
            'barcode': m['barcode'], 'subclass': m['subclass'], 'material': m['material'],
            'pack_format': m['pack_format'], 'factory': m['factory'],
            'factory_id': m['facid'], 'cost': m['cost'], 'cost_basis': m['cost_basis'],
            'retail': m['retail'], 'case_pack': case, 'inner': m['inner'],
            'hts': m['hts'], 'coo': m['coo'], 'port': m['port'],
            'plan': plan, 'direct': dqty, 'embedded': eqty, 'total': total,
            'diff': round(total - plan),
            'pos': direct.get(d, {}), 'pos_embedded': embedded.get(d, {}),
        })

    for d, m in all_master.items():
        if d not in ordered:
            findings['not_ordered'].append({'dpci': d, 'program': m['program'],
                                            'desc': str(m['desc'])[:45],
                                            'plan': round(m['plan'] or 0)})

    for k in ('cost_mismatch', 'retail_mismatch'):
        seen, uniq = set(), []
        for f in findings[k]:
            sig = (f['dpci'], f.get('po_unit') or f.get('po_resale'))
            if sig not in seen:
                seen.add(sig)
                uniq.append(f)
        findings[k] = uniq

    json.dump({'rows': rows, 'findings': findings, 'superseded': superseded,
               'assortments': asst, 'excluded_pos': sorted(excluded),
               'assortment_boxes': {k: v for k, v in ast_boxes.items()}},
              open(os.path.join(a.outdir, 'recon.json'), 'w', encoding='utf-8'),
              ensure_ascii=False, indent=1, default=str)

    print(json.dumps({
        'pos_in': len(pos), 'pos_kept': len(kept), 'pos_superseded': len(superseded),
        'superseded_with_qty_change': sum(1 for s in superseded if s['qty_changed']),
        'version_date_source': {k: sum(1 for h in kept if h['_dtsrc'] == k)
                                for k in ('stamp', 'filename', 'folder', 'none')},
        'dpci_ordered': len(ordered), 'dpci_reconciled': len(rows),
        'assortments_in_po': len(ast_boxes),
        'findings_counts': {k: len(v) for k, v in findings.items()},
    }, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
