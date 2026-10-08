#!/usr/bin/env python3
"""
build_grid.py — write the PO GRID workbook, one sheet per factory.

Layout follows the account team's hand-built GRIDs; see
references/grid-layout.md for the verified row/column map.

Photos are floating pictures with a TWO-CELL anchor ("move and size with
cells"), so they travel with their row when the sheet is sorted or filtered.

USAGE
  python build_grid.py --recon work/recon.json \
      --program "Decor & Night Of" \
      --images products_images.zip --images more_images.zip \
      --title "D240 26C5 HWN" \
      --out "D240_26C5_HWN_PO_GRID_Decor.xlsx" \
      --report build_report.json
"""
import argparse, json, os, re, subprocess, sys, tempfile, warnings, zipfile
from collections import defaultdict

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image as PILImage

warnings.filterwarnings('ignore')

# DC (Order # third segment) -> destination shown on row 5.
# ONLY add a code here once it has been confirmed against a hand-built GRID or
# by the account team. An unmapped DC falls through to its raw code and is
# reported as a header problem — that is deliberate. A wrong city name is worse
# than a visible code, because nothing downstream can tell it is wrong.
# DC (Order # third segment) -> destination shown on row 5.
# Only add a code the account team has confirmed. A wrong destination prints
# cleanly across every column of that DC and looks entirely normal, so an
# unmapped code is written raw and reported instead of guessed.
DEST = {
    '0581': 'LAS',         # confirmed: 雄健 and 從盛 reference GRIDs
    '3891': 'SAVANNAH',    # confirmed: 雄健 and 從盛 reference GRIDs
    '3758': 'CHARLESTON',  # confirmed by the account team, 2026-07-28
    '3887': 'HOUSTON',     # confirmed by the account team, 2026-07-28
    '3890': 'PNW',         # confirmed by Chris, 2026-08-18 (27C2 Easter)
}

# PO-type shading, keyed on the PO-type label in row 1, NOT on whether the PO
# carries an assortment. The team's own GRIDs shade only TS and T1; every type
# is shaded here so no column is left visually ungrouped. Low-saturation tints
# throughout, kept clear of the red used for the S/W band.
TYPE_FILL = {
    'TS': 'FFF2CC',      # cream
    'T1': 'E2EFDA',      # green
    'T2': 'DDEBF7',      # blue
    'T3': 'E4DFEC',      # purple
    'H1': 'D9E1F2',      # steel blue
    'H3': 'EDEDED',      # grey
    # Mini Seasonal programs use a different tranche vocabulary entirely
    # (from "Assigned by transaction set sender:" on the PO, not the
    # filename) — Basic / Seasonal / Set Order. The account team's own
    # reference GRID for 27C1 leaves these unfilled, but the skill shades
    # every type so no PO column reads as visually ungrouped.
    'BAS': 'FCE4D6',      # peach
    'SEA': 'D9D2E9',      # lavender
    'SET': 'D0E0E3',      # teal
    # 27C2 Easter adds two more codes from the same 'Assigned by
    # transaction set sender:' field — New Item and Special Order.
    'NWT': 'FFF2CC',     # cream  (New Item)
    'SPO': 'E2EFDA',     # green  (Special Order)
}
SW_FILL = 'FADBD8'       # shipping-window band, light red

FIELDS = ['DPCI', 'STYLE', 'PICTURE', 'DESCRIPTION', 'UPC', 'MATERIAL',
          'AGE', 'MAKER', 'Retail$/Packaging Type', 'IN', 'CTN', 'assortment']
WIDTHS = [13.5, 20.75, 17.625, 32.25, 17.75, 15.375, 5.75, 12.625, 12.875,
          4.375, 4.375, 6.5, 11.0]
C_ASST_UNITS = 12                       # the narrow "assortment" column (L)
C_ROWLBL = 13                           # names what rows 3/4/5 hold (M)
FIRST_PO_COL = 14                       # PO columns start at N
ROW_LABELS = {3: 'TG PO#', 4: 'S/W', 5: 'DES PORT'}

THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill('solid', fgColor='D9D9D9')
YELLOW = PatternFill('solid', fgColor='FFFF00')
FONT = 'Arial'
ROW_H, ASST_ROW_H, GAP_ROW_H = 78, 34.5, 18.75
IMG_BOX = 96                            # px, longest edge
COL_C_PX, ROW_PX = 129, 104


def short_name(f):
    s = re.sub(r'\b(CO|LTD|LIMITED|COMPANY|INC|CORP)\b\.?,?', '', str(f), flags=re.I)
    s = re.sub(r'[.,]+', ' ', s)
    return (re.sub(r'\s+', ' ', s).strip() or str(f))[:28]


def maker_name(f):
    """Factory name for the MAKER cell: legal suffixes stripped, NOT cut to 28
    characters (that limit only exists for sheet tabs)."""
    s = re.sub(r'\b(CO|LTD|LIMITED|COMPANY|INC|CORP)\b\.?,?', '', str(f), flags=re.I)
    s = re.sub(r'[.,]+', ' ', s)
    return re.sub(r'\s+', ' ', s).strip() or str(f)


def sheet_title(name, used):
    t = re.sub(r'[\\/*?:\[\]]', '-', short_name(name))[:31] or 'SHEET'
    base, n = t, 2
    while t in used:
        t, n = '%s_%d' % (base[:28], n), n + 1
    used.add(t)
    return t


def load_images(specs, workdir):
    imgs = {}
    for spec in specs:
        if os.path.isdir(spec):
            for fn in os.listdir(spec):
                if fn.lower().endswith(('.png', '.jpg', '.jpeg')):
                    imgs[os.path.splitext(fn)[0]] = os.path.join(spec, fn)
            continue
        with zipfile.ZipFile(spec) as z:
            for n in z.namelist():
                if not n.lower().endswith(('.png', '.jpg', '.jpeg')):
                    continue
                dp = os.path.splitext(os.path.basename(n))[0]
                p = os.path.join(workdir, os.path.basename(n))
                open(p, 'wb').write(z.read(n))
                imgs[dp] = p
    return imgs


def ship_key(meta):
    """PO columns run left to right by shipping start date."""
    s = meta.get('ship_start') or ''
    return (s[-4:] + s[:5]) if len(s) == 10 else 'zzzz'


def place_image(ws, row, path, report, dpci):
    try:
        with PILImage.open(path) as im:
            w0, h0 = im.size
        sc = min(IMG_BOX / w0, IMG_BOX / h0)
        w, h = int(w0 * sc), int(h0 * sc)
        img = XLImage(path)
        ox, oy = max((COL_C_PX - w) // 2, 0), max((ROW_PX - h) // 2, 0)
        # Two-cell anchor: both markers sit in this row, so the picture moves,
        # sizes and collapses with it. A one-cell anchor does none of that.
        img.anchor = TwoCellAnchor(
            editAs='twoCell',
            _from=AnchorMarker(col=2, colOff=pixels_to_EMU(ox),
                               row=row - 1, rowOff=pixels_to_EMU(oy)),
            to=AnchorMarker(col=2, colOff=pixels_to_EMU(ox + w),
                            row=row - 1, rowOff=pixels_to_EMU(oy + h)))
        ws.add_image(img)
        return True
    except Exception:                                            # noqa: BLE001
        report['no_image'].append(dpci)
        return False


def write_summary_from_json(wb, data):
    """First tab written from a caller-supplied summary (e.g. the TG Team PO
    platform's Chinese summary), so the GRID and the validation report show
    the same checks in the same words. Layout: title, run line, headline
    figures, the check table (OK green / REVIEW orange), then one detail
    table per non-empty finding."""
    ws = wb.create_sheet('Validation Summary', 0)
    grey = PatternFill('solid', fgColor='D9D9D9')
    green = PatternFill('solid', fgColor='E2EFDA')
    orange = PatternFill('solid', fgColor='F8CBAD')
    money = set(data.get('money_cols', []))
    qty = set(data.get('qty_cols', []))
    ws['A1'] = data.get('title', '')
    ws['A1'].font = Font(FONT, 14, bold=True)
    ws['A2'] = 'PO Validation Summary'
    ws['A2'].font = Font(FONT, 12, bold=True, italic=True)
    if data.get('run'):
        ws['A3'] = data['run']
        ws['A3'].font = Font(FONT, 9, color='808080')
    r = 5
    for k, v in data.get('headline', []):
        ws.cell(r, 1, k).font = Font(FONT, 10, bold=True)
        c = ws.cell(r, 2, v)
        c.alignment = Alignment(horizontal='left')
        c.font = Font(FONT, 10)
        if k == '結論':
            c.font = Font(FONT, 10, bold=True, color='C00000' if '需確認' in str(v) else '548235')
        r += 1
    r += 1
    for j, h in enumerate(['Check 檢核項目', 'Count', 'Status'], 1):
        c = ws.cell(r, j, h)
        c.font, c.fill = Font(FONT, 10, bold=True), grey
    r += 1
    for label, n, status in data.get('checks', []):
        fill = orange if status == 'REVIEW' else (green if status == 'OK' else None)
        for j, v in enumerate([label, n, status], 1):
            c = ws.cell(r, j, v)
            c.font = Font(FONT, 10)
            if fill:
                c.fill = fill
        r += 1
    for d in data.get('details', []):
        r += 1
        ws.cell(r, 1, d['title']).font = Font(FONT, 11, bold=True)
        r += 1
        cols = d['columns']
        for j, h in enumerate(cols, 1):
            c = ws.cell(r, j, h)
            c.font, c.fill = Font(FONT, 10, bold=True), grey
        r += 1
        for row in d['rows']:
            for j, v in enumerate(row, 1):
                c = ws.cell(r, j, v)
                c.font = Font(FONT, 10)
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    if cols[j - 1] in money:
                        c.number_format = '"$"#,##0.00##'
                    elif cols[j - 1] in qty:
                        c.number_format = '#,##0'
                    elif cols[j - 1] == '差異 %':
                        c.number_format = '0.0"%"'
            r += 1
    ws.column_dimensions['A'].width = 54
    for col, w in zip('BCDEFGH', [26, 22, 22, 22, 40, 14, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A5'


def write_validation_summary(wb, recon, title, gaps):
    """First tab: the PO-vs-master cross-check results, so a mismatch is
    visible in the workbook itself and not just in a chat message or a JSON
    report file that gets separated from the GRID. Standing instruction:
    always surface a PO-vs-master inconsistency rather than silently picking
    a side (see po-validation-workflow SOP)."""
    ws = wb.create_sheet('Validation Summary', 0)
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 60
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 20
    bold = Font(FONT, 11, bold=True)
    hdr_fill = PatternFill('solid', fgColor='D9D9D9')
    red_fill = PatternFill('solid', fgColor='F8CBAD')
    amber_fill = PatternFill('solid', fgColor='FFF2CC')
    green_fill = PatternFill('solid', fgColor='E2EFDA')

    ws.cell(1, 1, title).font = Font(FONT, 14, bold=True)
    ws.cell(2, 1, 'PO Validation Summary').font = Font(FONT, 12, bold=True, italic=True)

    findings = recon.get('findings', {})
    superseded = recon.get('superseded', [])
    labels = [
        ('unknown_dpci', 'Unknown DPCI (PO item not in master)'),
        ('cost_mismatch', 'Cost mismatch (PO unit cost vs master)'),
        ('retail_mismatch', 'Retail mismatch (PO resale vs master)'),
        ('qty_vs_plan', "Qty vs plan (beyond whole-case rounding)"),
        ('assortment_mismatch', 'Assortment mismatch (box vs component qty)'),
        ('not_ordered', 'Master DPCI not yet ordered'),
        ('duplicate_po_groups', 'Duplicate PO groups (same content, diff PO#)'),
    ]
    r = 4
    for c, h in ((1, 'Check'), (2, 'Count'), (3, 'Status')):
        cell = ws.cell(r, c, h)
        cell.font, cell.fill = bold, hdr_fill
    r += 1
    for key, label in labels:
        n = len(findings.get(key, []))
        ws.cell(r, 1, label)
        ws.cell(r, 2, n)
        ws.cell(r, 3, 'OK' if n == 0 else 'REVIEW')
        fill = green_fill if n == 0 else red_fill
        for c in (1, 2, 3):
            ws.cell(r, c).fill = fill
        r += 1
    ws.cell(r, 1, 'PO versions superseded (revision applied)')
    ws.cell(r, 2, len(superseded))
    for c in (1, 2):
        ws.cell(r, c).fill = amber_fill if superseded else green_fill
    r += 1
    ws.cell(r, 1, 'Header gaps (PO type / ship window missing on this template)')
    ws.cell(r, 2, len(gaps))
    for c in (1, 2):
        ws.cell(r, c).fill = red_fill if gaps else green_fill
    r += 2

    for key, label in labels:
        rows_ = findings.get(key, [])
        if not rows_:
            continue
        ws.cell(r, 1, label + ' — detail').font = bold
        r += 1
        cols = sorted({k for row in rows_ for k in row.keys()})
        for i, c in enumerate(cols, start=1):
            ws.cell(r, i, c).font = Font(FONT, 9, bold=True)
        r += 1
        for row in rows_:
            for i, c in enumerate(cols, start=1):
                v = row.get(c)
                # some findings carry a list (e.g. unknown_dpci -> the POs that
                # ordered it); openpyxl cannot write a list to a cell
                if isinstance(v, (list, tuple, set)):
                    v = ', '.join(str(x) for x in v)
                elif isinstance(v, dict):
                    v = ', '.join('%s=%s' % (k, x) for k, x in v.items())
                ws.cell(r, i, v)
            r += 1
        r += 1

    if gaps:
        ws.cell(r, 1, 'Header gap detail (%d POs)' % len(gaps)).font = bold
        r += 1
        for g in gaps:
            ws.cell(r, 1, g.get('po'))
            ws.cell(r, 2, g.get('missing'))
            ws.cell(r, 3, g.get('file'))
            r += 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--recon', required=True)
    ap.add_argument('--items')
    ap.add_argument('--program')
    ap.add_argument('--images', action='append', default=[])
    ap.add_argument('--title', default='PO GRID')
    ap.add_argument('--out', required=True)
    ap.add_argument('--report', default='build_report.json')
    ap.add_argument('--no-recalc', action='store_true')
    ap.add_argument('--summary-json',
                    help='write the first tab from this summary JSON instead of '
                         'the built-in English recon summary')
    ap.add_argument('--allow-header-gaps', action='store_true',
                    help='build even if a PO column lacks a type label '
                         'or shipping window (normally an error)')
    a = ap.parse_args()

    recon = json.load(open(a.recon, encoding='utf-8'))
    asst_def = recon.get('assortments', {})
    asst_boxes = recon.get('assortment_boxes', {})
    comp_of = defaultdict(list)                  # component DPCI -> [(asst, units)]
    for adpci, spec in asst_def.items():
        # A proposal-only group whose box DPCI never actually showed up as an
        # AST line on a real PO (e.g. a "TBD" placeholder, or components the
        # account team decided to order individually instead) is not a box
        # that was purchased — its components go in the GRID as ordinary
        # standalone items, not under a fabricated assortment header.
        if adpci not in asst_boxes:
            continue
        for c in spec.get('comp', []):
            comp_of[c['dpci']].append((adpci, c['units']))

    items_path = a.items or os.path.join(os.path.dirname(a.recon), 'items.json')
    items = json.load(open(items_path, encoding='utf-8'))

    # UPC of each assortment master, straight off the PO
    asst_upc = {}
    for i in items['items']:
        if i['kind'] == 'ast':
            asst_upc['%s-%s-%s' % (i['sku'][:3], i['sku'][3:5], i['sku'][5:])] = i.get('upc')

    def ver(h):
        s = h.get('retrieved') or ''
        m = re.match(r'(\d{4})/(\d{1,2})/(\d{1,2})\s+(\d{1,2}):(\d{2})', s)
        if m:
            return tuple(int(x) for x in m.groups())
        fd = h.get('fname_date') or ''
        return (0, int(fd[:2]), int(fd[2:]), 0, 0) if len(fd) == 4 else (0, 0, 0, 0, 0)

    heads = defaultdict(list)
    for h in items['pos']:
        heads[h['po']].append(h)
    po_meta = {}
    for po, hs in heads.items():
        win = max(hs, key=ver)
        if not win.get('dclabel'):
            for alt in hs:
                if alt.get('dclabel'):
                    win = dict(win, dclabel=alt['dclabel'])
                    break
        if not (win.get('ship_start') and win.get('ship_end')):
            for alt in sorted(hs, key=ver, reverse=True):
                if alt.get('ship_start') and alt.get('ship_end'):
                    win = dict(win, ship_start=alt['ship_start'],
                               ship_end=alt['ship_end'])
                    break
        po_meta[po] = win

    workdir = tempfile.mkdtemp()
    imgs = load_images(a.images, workdir)

    rows = [r for r in recon['rows'] if not a.program or r['program'] == a.program]
    if not rows:
        raise SystemExit('no rows for program %r' % a.program)

    # Header gate. The line-item gates in parse_po_pdfs.py check quantities and
    # money; nothing checked the PO header fields, so a regex that stopped
    # matching left blank cells and shipped quietly. Every PO column must carry
    # a type label and a shipping window, or the build stops.
    used_pos = set()
    for r in rows:
        used_pos |= set(r['pos']) | set(r['pos_embedded'])
    gaps = []
    for po in sorted(used_pos):
        m = po_meta.get(po)
        if not m:
            gaps.append({'po': po, 'missing': 'no parsed PO header at all'})
            continue
        miss = []
        if not m.get('dclabel'):
            miss.append('PO type')
        if not (m.get('ship_start') and m.get('ship_end')):
            miss.append('shipping window')
        if miss:
            gaps.append({'po': po, 'missing': ' + '.join(miss),
                         'file': m.get('file')})
    if gaps and not a.allow_header_gaps:
        print(json.dumps({
            'error': 'PO header fields missing — refusing to build',
            'gaps': gaps,
            'hint': 'A parser regex has probably stopped matching. Check '
                    'ship_start / ship_end and dclabel in items.json for these '
                    'POs, or pass --allow-header-gaps to build anyway.',
        }, ensure_ascii=False, indent=2))
        raise SystemExit(1)
    fac = defaultdict(list)
    for r in rows:
        fac[r['factory'] or '(UNASSIGNED)'].append(r)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    if a.summary_json:
        write_summary_from_json(wb, json.load(open(a.summary_json, encoding='utf-8')))
    else:
        write_validation_summary(wb, recon, a.title, gaps)
    used, report = set(), {'sheets': [], 'no_image': [], 'no_po_meta': [],
                           'header_gaps': gaps, 'unmapped_dc': {}}

    for factory in sorted(fac, key=lambda f: (-len(fac[f]), str(f))):
        recs = {r['dpci']: r for r in sorted(fac[factory], key=lambda r: r['dpci'])}
        pos = set()
        for r in recs.values():
            pos |= set(r['pos']) | set(r['pos_embedded'])
        for p in sorted(pos - set(po_meta)):
            report['no_po_meta'].append(p)
        pos = sorted([p for p in pos if p in po_meta],
                     key=lambda p: (ship_key(po_meta[p]), p))

        ws = wb.create_sheet(sheet_title(factory, used))
        ncol = FIRST_PO_COL - 1 + len(pos)
        c_total, c_commit = ncol + 1, ncol + 2

        # assortments whose components sit on this sheet — only boxes actually
        # confirmed by a real PO AST line (see comp_of above); a proposal-only
        # group never gets a header row here either
        sheet_asst = [ad for ad, spec in asst_def.items()
                      if ad in asst_boxes
                      and any(c['dpci'] in recs for c in spec.get('comp', []))]

        for i, w in enumerate(WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for i in range(FIRST_PO_COL, ncol + 1):
            ws.column_dimensions[get_column_letter(i)].width = 17.375
        ws.column_dimensions[get_column_letter(c_total)].width = 17
        ws.column_dimensions[get_column_letter(c_commit)].width = 18.125
        for r, h in ((1, 24.75), (2, 35.25), (3, 19.5), (4, 18.75), (5, 18.75)):
            ws.row_dimensions[r].height = h

        ws.cell(1, 1, a.title).font = Font(FONT, 14, bold=True)
        for i, name in enumerate(FIELDS, start=1):
            c = ws.cell(2, i, name)
            c.font = Font(FONT, 10, bold=True)
            c.fill = YELLOW if i == C_ASST_UNITS else HDR_FILL
            c.border = BORDER
            c.alignment = Alignment('center', 'center', wrap_text=True)

        for row, lbl in ROW_LABELS.items():
            c = ws.cell(row, C_ROWLBL, lbl)
            c.font = Font(FONT, 9, bold=True)
            c.alignment = Alignment('right', 'center')
            c.border = BORDER
            if row == 4:                       # S/W label joins its own band
                c.fill = PatternFill('solid', fgColor=SW_FILL)

        po_fill = {}
        for i, po in enumerate(pos):
            col, m = FIRST_PO_COL + i, po_meta[po]
            lbl = m.get('dclabel', '')
            # PO type and shipping window are hard-gated before any sheet is
            # written. An unmapped DC is different: it is not a parse failure,
            # it is a code nobody has confirmed a destination for yet. Show the
            # raw code so it is visible on the sheet, and surface it in the
            # report rather than blocking the build.
            if m.get('dc') not in DEST:
                report['unmapped_dc'].setdefault(m.get('dc'), []).append(po)
            fillc = TYPE_FILL.get(lbl)
            po_fill[po] = PatternFill('solid', fgColor=fillc) if fillc else None
            for row, val, fnt in ((1, lbl, Font(FONT, 10, bold=True)),
                                  (3, po, Font(FONT, 10, bold=True)),
                                  (4, '%s-%s' % (m['ship_start'][:5], m['ship_end'][:5])
                                      if m.get('ship_start') and m.get('ship_end') else '',
                                      Font(FONT, 9, bold=True)),
                                  (5, DEST.get(m.get('dc'), m.get('dc')), Font(FONT, 9))):
                c = ws.cell(row, col, val)
                c.font, c.alignment = fnt, Alignment('center', 'center')
                if po_fill[po] and row in (1, 3):
                    c.fill = po_fill[po]
                if row == 4:
                    c.fill = PatternFill('solid', fgColor=SW_FILL)
                if row in (3, 4):
                    c.border = BORDER
            ws.cell(2, col).fill, ws.cell(2, col).border = HDR_FILL, BORDER

        for col, lbl in ((c_total, ' PO TOTAL'), (c_commit, "100% COMMIT\nQ'TY")):
            c = ws.cell(2, col, lbl)
            c.font, c.fill, c.border = Font(FONT, 10, bold=True), HDR_FILL, BORDER
            c.alignment = Alignment('center', 'center', wrap_text=True)

        def style_row(r, last_col=None):
            for col in range(1, (last_col or c_commit) + 1):
                c = ws.cell(r, col)
                c.border, c.font = BORDER, Font(FONT, 10)
                if col == 4:
                    c.alignment = Alignment('left', 'center', wrap_text=True)
                elif col in (1, 2, 6, 7, 8, 9):
                    c.alignment = Alignment('center', 'center', wrap_text=True)
                else:
                    c.alignment = Alignment('center', 'center')
                if col >= FIRST_PO_COL:
                    c.number_format = '#,##0'


        def write_item(r, rec):
            ws.row_dimensions[r].height = ROW_H
            ws.cell(r, 1, rec['dpci'])
            ws.cell(r, 2, rec['style'])
            ws.cell(r, 4, rec['desc'])
            ws.cell(r, 5, str(rec['barcode'] or ''))
            ws.cell(r, 6, rec['material'])
            # factory name and ID on separate lines so the column can stay narrow
            ws.cell(r, 8, '%s\n%s' % (maker_name(factory), rec['factory_id'] or '')
                    if rec['factory_id'] else maker_name(factory))
            retail = rec['retail']
            seal = ('$%.2f' % retail) if isinstance(retail, (int, float)) else ''
            if rec.get('pack_format'):
                seal += '\n%s' % rec['pack_format']
            ws.cell(r, 9, seal)
            ws.cell(r, 10, rec['inner'])
            ws.cell(r, 11, rec['case_pack'])
            if comp_of.get(rec['dpci']):
                ws.cell(r, C_ASST_UNITS, comp_of[rec['dpci']][0][1])
            for i, po in enumerate(pos):
                q = rec['pos'].get(po, 0) + rec['pos_embedded'].get(po, 0)
                if q:
                    ws.cell(r, FIRST_PO_COL + i, q)
            ws.cell(r, c_total, '=SUM(%s%d:%s%d)'
                    % (get_column_letter(FIRST_PO_COL), r, get_column_letter(ncol), r))
            ws.cell(r, c_commit, rec['plan'])
            style_row(r)
            if comp_of.get(rec['dpci']):
                ws.cell(r, C_ASST_UNITS).fill = YELLOW
            p = imgs.get(rec['dpci'])
            if p:
                place_image(ws, r, p, report, rec['dpci'])
            else:
                report['no_image'].append(rec['dpci'])

        r = 6
        written = set()
        # assortment groups first: a header row carrying the box counts in the
        # PO columns they were ordered on, then that box's component items
        for adpci in sorted(sheet_asst):
            spec = asst_def[adpci]
            boxes = asst_boxes.get(adpci, {})
            ws.row_dimensions[r].height = ASST_ROW_H
            ws.cell(r, 1, adpci)
            ws.cell(r, 2, 'ASSORTMENT-%s' % (spec.get('style') or ''))
            ws.cell(r, 4, spec.get('desc'))
            ws.cell(r, 5, str(asst_upc.get(adpci) or ''))
            for i, po in enumerate(pos):
                if boxes.get(po):
                    c = ws.cell(r, FIRST_PO_COL + i, '%s\n(%d)' % (adpci, boxes[po]))
                    c.fill = YELLOW
            style_row(r)
            ws.cell(r, 1).fill = YELLOW
            ws.cell(r, 2).font = Font(FONT, 10, bold=True)
            for i, po in enumerate(pos):
                ws.cell(r, FIRST_PO_COL + i).alignment = Alignment(
                    'center', 'center', wrap_text=True)
            r += 1
            for c in spec.get('comp', []):
                if c['dpci'] in recs and c['dpci'] not in written:
                    write_item(r, recs[c['dpci']])
                    written.add(c['dpci'])
                    r += 1
            ws.row_dimensions[r].height = GAP_ROW_H
            r += 1

        for dpci, rec in recs.items():
            if dpci not in written:
                write_item(r, rec)
                r += 1

        # MAKER (H) width follows the longest factory name on this sheet
        # (very long names wrap onto a second line inside the 78pt-high row)
        maker_w = min(len(maker_name(factory)), 30) * 1.1 + 2
        ws.column_dimensions['H'].width = round(max(maker_w, 12.625), 2)
        ws.freeze_panes = 'D6'
        ws.sheet_view.zoomScale = 100
        report['sheets'].append({'sheet': ws.title, 'factory': factory,
                                 'items': len(recs), 'po_cols': len(pos),
                                 'assortments': sorted(sheet_asst)})

    wb.save(a.out)

    if not a.no_recalc:
        rc = '/mnt/skills/public/xlsx/scripts/recalc.py'
        if os.path.exists(rc):
            try:
                out = subprocess.run([sys.executable, rc, a.out, '120'],
                                     capture_output=True, text=True, timeout=600)
                report['recalc'] = json.loads(
                    out.stdout[out.stdout.index('{'):out.stdout.rindex('}') + 1])
            except Exception as e:                               # noqa: BLE001
                report['recalc'] = {'error': repr(e)}

    report['file'] = a.out
    report['sheet_count'] = len(report['sheets'])
    report['item_count'] = sum(s['items'] for s in report['sheets'])
    report['image_count'] = report['item_count'] - len(report['no_image'])
    json.dump(report, open(a.report, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(json.dumps({k: report[k] for k in
                      ('file', 'sheet_count', 'item_count', 'image_count',
                       'no_image', 'no_po_meta', 'header_gaps',
                       'unmapped_dc', 'recalc') if k in report},
                     ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
