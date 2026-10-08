#!/usr/bin/env python3
"""
verify_layout.py — check a PO GRID against the frozen layout contract.

The layout is identical for every program and every season. This script proves
it mechanically, so a GRID built in one session can be checked against the spec
without eyeballing it beside last season's file.

Run it on anything claiming to be a PO GRID — including a workbook produced by
a different session or rebuilt after `scripts/` went missing.

USAGE
  python verify_layout.py GRID.xlsx
  python verify_layout.py GRID.xlsx --quiet     # exit code only

Exit 0 = conforms. Exit 1 = does not.
"""
import argparse, json, sys, warnings

import openpyxl

warnings.filterwarnings('ignore')

HEADERS = ['DPCI', 'STYLE', 'PICTURE', 'DESCRIPTION', 'UPC', 'MATERIAL', 'AGE',
           'MAKER', 'Retail$/Packaging Type', 'IN', 'CTN', 'assortment']
ROW_LABELS = {3: 'TG PO#', 4: 'S/W', 5: 'DES PORT'}
C_LABEL = 13                      # M
FIRST_PO = 14                     # N
FIRST_DATA_ROW = 6
TYPE_FILL = {'TS': 'FFF2CC', 'T1': 'E2EFDA', 'T2': 'DDEBF7',
             'T3': 'E4DFEC', 'H1': 'D9E1F2', 'H3': 'EDEDED',
             'BAS': 'FCE4D6', 'SEA': 'D9D2E9', 'SET': 'D0E0E3',
             'NWT': 'FFF2CC', 'SPO': 'E2EFDA'}
SW_FILL = 'FADBD8'
ASST_FILL = 'FFFF00'


def fill_of(cell):
    try:
        if cell.fill.fill_type != 'solid':
            return None
        rgb = cell.fill.start_color.rgb
        return rgb[-6:] if isinstance(rgb, str) else None
    except Exception:                                            # noqa: BLE001
        return None


def check(path):
    wb = openpyxl.load_workbook(path)
    problems = []

    def bad(sheet, what):
        problems.append({'sheet': sheet, 'problem': what})

    for ws in wb.worksheets:
        t = ws.title
        if t == 'Validation Summary':
            continue

        for i, h in enumerate(HEADERS, start=1):
            got = ws.cell(2, i).value
            if str(got or '').strip() != h:
                bad(t, 'col %d header is %r, expected %r' % (i, got, h))

        for row, lbl in ROW_LABELS.items():
            got = ws.cell(row, C_LABEL).value
            if str(got or '').strip() != lbl:
                bad(t, 'M%d is %r, expected %r' % (row, got, lbl))

        if ws.freeze_panes != 'D6':
            bad(t, 'freeze panes %r, expected D6' % ws.freeze_panes)

        first = None
        for r in range(1, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if v and str(v).strip().count('-') == 2:
                first = r
                break
        if first is not None and first != FIRST_DATA_ROW:
            bad(t, 'first data row is %d, expected %d' % (first, FIRST_DATA_ROW))

        po_cols = [c for c in range(FIRST_PO, ws.max_column + 1)
                   if ws.cell(3, c).value]
        if not po_cols:
            bad(t, 'no PO columns found from column N onward')
            continue
        if min(po_cols) != FIRST_PO:
            bad(t, 'PO columns start at %d, expected %d' % (min(po_cols), FIRST_PO))

        for c in po_cols:
            lbl = str(ws.cell(1, c).value or '').strip()
            po = ws.cell(3, c).value
            if not lbl:
                bad(t, 'PO %s has no type label in row 1' % po)
            elif lbl not in TYPE_FILL:
                bad(t, 'PO %s type %r has no defined fill' % (po, lbl))
            else:
                for r in (1, 3):
                    if fill_of(ws.cell(r, c)) != TYPE_FILL[lbl]:
                        bad(t, 'PO %s row %d fill %s, expected %s'
                            % (po, r, fill_of(ws.cell(r, c)), TYPE_FILL[lbl]))
            if not ws.cell(4, c).value:
                bad(t, 'PO %s has no shipping window' % po)
            elif not ws.cell(4, c).font.bold:
                bad(t, 'PO %s shipping window is not bold' % po)
            if fill_of(ws.cell(4, c)) != SW_FILL:
                bad(t, 'PO %s S/W fill %s, expected %s'
                    % (po, fill_of(ws.cell(4, c)), SW_FILL))
            if not ws.cell(5, c).value:
                bad(t, 'PO %s has no destination' % po)

        if fill_of(ws.cell(4, C_LABEL)) != SW_FILL:
            bad(t, 'S/W label cell is not on the red band')
        if fill_of(ws.cell(2, 12)) != ASST_FILL:
            bad(t, 'assortment header (L2) is not yellow')

        last = max(po_cols)
        if str(ws.cell(2, last + 1).value or '').strip() != 'PO TOTAL':
            bad(t, 'expected " PO TOTAL" after the PO columns, got %r'
                % ws.cell(2, last + 1).value)
        if "COMMIT" not in str(ws.cell(2, last + 2).value or ''):
            bad(t, "expected the 100%% COMMIT column last, got %r"
                % ws.cell(2, last + 2).value)

        for r in range(FIRST_DATA_ROW, ws.max_row + 1):
            b = ws.cell(r, 2).value
            if b and str(b).startswith('ASSORTMENT-'):
                if fill_of(ws.cell(r, 1)) != ASST_FILL:
                    bad(t, 'row %d assortment DPCI is not yellow' % r)
                if not ws.cell(r, 5).value:
                    bad(t, 'row %d assortment has no UPC' % r)

    return wb, problems


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('grid')
    ap.add_argument('--quiet', action='store_true')
    a = ap.parse_args()

    wb, problems = check(a.grid)
    out = {'file': a.grid, 'sheets': len(wb.sheetnames),
           'conforms': not problems, 'problem_count': len(problems),
           'problems': problems[:40]}
    if not a.quiet:
        print(json.dumps(out, ensure_ascii=False, indent=2))
    return 0 if not problems else 1


if __name__ == '__main__':
    sys.exit(main())
